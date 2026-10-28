"""Builds the midterm workbook and the BTC price file it reads.
Run:  uv run --with yfinance --with openpyxl python build_sheet.py
      → writes data/btc_2025_2026.csv and data/midterm_sheet.xlsx next to this file.

The workbook carries no numbers of ours except the BTC closes: every other cell is
the student's, or a formula on the student's cells. Google Sheets ▸ File ▸ Import
keeps the tabs and the formulas."""
from datetime import date
from pathlib import Path

import yfinance as yf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).parent
DATA = ROOT / "data"
START, END = "2025-01-01", "2026-10-26"  # END exclusive: closes up to 25 Oct 2026

GOLD = PatternFill("solid", fgColor="CDBA80")
INK = Font(bold=True, color="363636")
RED = Font(bold=True, color="AF1F25")
WRAP = Alignment(wrap_text=True, vertical="top")

# ---- BTC daily closes -------------------------------------------------------
px = yf.download("BTC-USD", start=START, end=END, progress=False, auto_adjust=False)
close = px["Close"].squeeze().round(4)
close.index = close.index.date
close.rename("close").rename_axis("Date").to_csv(DATA / "btc_2025_2026.csv")


def ref(x):
    """A link to another cell that stays empty while that cell is empty."""
    return f'=IF({x}="","",{x})'


def sheet(wb, title, header, rows, widths, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = RED
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header))
        ws.row_dimensions[1].height = 45
        ws.cell(1, 1).alignment = WRAP
        r0 = 3
    for j, h in enumerate(header, 1):
        c = ws.cell(r0, j, h)
        c.fill, c.font, c.alignment = GOLD, INK, WRAP
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v).alignment = WRAP
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    return ws


wb = Workbook()
ws = wb.active
ws.title = "Start"
start = [
    "Midterm — Coinbase's third quarter, read by you and an AI",
    "AI for Business and FinTech · ESE Florence · 40% of the final mark",
    "",
    "One tab per step. Fill them in order; the brief is on the course site (week 5 ▸ Midterm).",
    "1 Model — the model you use, read from its model card.",
    "2 Look first — Coinbase's last five quarters next to the BTC price, before any prediction.",
    "3 Preview — five predictions for Q3 2026. Fill BEFORE Thursday 29 October, 21:00 Italian time.",
    "4 Test — ten questions on the Q3 document with answers found by hand, three traps, the temperature test.",
    "5 Cost — what one answer and one analyst note cost, and what they are worth.",
    "6 Score — your predictions against the actual Q3 numbers and against 'Q3 = Q2'.",
    "BTC daily — Bitcoin daily closes in USD, 1 Jan 2025 – 25 Oct 2026 (Yahoo Finance, BTC-USD). Do not edit.",
    "",
    "Grey cells with a formula fill themselves. Every other number is yours: write where you found it.",
    "Upload: File ▸ Download ▸ Microsoft Excel (.xlsx), to Turnitin Midterm Submission by Tuesday 3 November, 23:59.",
]
for i, line in enumerate(start, 1):
    ws.cell(i, 1, line).font = RED if i == 1 else (INK if i in (7,) else Font(color="363636"))
ws.column_dimensions["A"].width = 120

# ---- 1 Model ----------------------------------------------------------------
sheet(wb, "1 Model",
      ["Item", "My answer", "Where I read it (link, section)"],
      [["Model I use in AI Studio", None, None],
       ["Model card or release page", None, None],
       ["Context window (input tokens)", None, None],
       ["Longest answer (output tokens)", None, None],
       ["Knowledge cutoff", None, None],
       ["Input price per 1M tokens ($)", None, None],
       ["Output price per 1M tokens ($)", None, None],
       ["Tokens in my documents (AI Studio token counter, all files attached)", None, None],
       ["Do they fit in the context window?", '=IF(OR(B4="",B9=""),"",IF(B9<=B4,"yes","no"))', "formula"],
       ["Which of my documents can the model have seen in training, and so what? (one line)", None, None]],
      [48, 34, 50])

# ---- 2 Look first -----------------------------------------------------------
quarters = [("Q2 2025", date(2025, 4, 1), date(2025, 6, 30)),
            ("Q3 2025", date(2025, 7, 1), date(2025, 9, 30)),
            ("Q4 2025", date(2025, 10, 1), date(2025, 12, 31)),
            ("Q1 2026", date(2026, 1, 1), date(2026, 3, 31)),
            ("Q2 2026", date(2026, 4, 1), date(2026, 6, 30)),
            ("Q3 2026", date(2026, 7, 1), date(2026, 9, 30))]
rows = []
for k, (q, a, b) in enumerate(quarters, 4):
    avg = f"=ROUND(AVERAGEIFS('BTC daily'!$B:$B,'BTC daily'!$A:$A,\">=\"&B{k},'BTC daily'!$A:$A,\"<=\"&C{k}),0)"
    later = "after 29 Oct" if q == "Q3 2026" else None
    rows.append([q, a, b, later, later, later, None, None, avg])
rows.append([None] * 9)
rows.append(["What I see (two lines):", None, None, None, None, None, None, None, None])
ws = sheet(wb, "2 Look first",
           ["Quarter", "From", "To", "Total revenue ($M)", "Transaction revenue ($M)",
            "Subscription & services revenue ($M)", "Where found (document, slide or page)",
            "Checked by hand? (yes/no)", "BTC average close ($)"],
           rows, [10, 12, 12, 14, 16, 18, 34, 14, 16],
           note="Fill the revenue columns from the documents in your notebook, then check each number by hand "
                "against the slide or page. Then: Insert ▸ Chart, revenue and BTC side by side. What moves with what?")
for k in range(4, 10):
    for col in "BC":
        ws[f"{col}{k}"].number_format = "yyyy-mm-dd"

# ---- 3 Preview --------------------------------------------------------------
naive = {"Total revenue ($M)": ref("'2 Look first'!D8"),
         "Transaction revenue ($M)": ref("'2 Look first'!E8"),
         "Subscription & services revenue ($M)": ref("'2 Look first'!F8")}
rows = [[f"P{i}", m, None, naive.get(m), None, None, None, None]
        for i, m in enumerate(list(naive) + ["(your choice)", "(your choice)"], 1)]
sheet(wb, "3 Preview",
      ["#", "What I predict for Q3 2026", "My number (one number; for a range, the midpoint)",
       "Naive: Q3 = Q2", "Coinbase's own Q3 outlook, if it gave one", "Where my number comes from (document, slide)",
       "Confidence (0–100)", "Why, in one line"],
      rows, [6, 34, 18, 14, 20, 34, 12, 40],
      note="Fill this tab BEFORE Thursday 29 October, 21:00 Italian time, when Coinbase publishes. "
           "Google Sheets keeps the version history (File ▸ Version history): that is the proof it came first.")

# ---- 4 Test -----------------------------------------------------------------
rows = [[f"G{i}", None, None, None, None, None, None, None, None] for i in range(1, 11)]
rows += [["T1", "No documents (AI Studio, no files, Google Search off): What was Coinbase's total revenue in Q3 2026?",
          None, None, None, None, None, None, None],
         ["T2", "(a question that needs two documents, e.g. the Q3 result against the Q2 outlook)",
          None, None, None, None, None, None, None],
         ["T3", "(your own trap; say in 'note' what it tests)", None, None, None, None, None, None, None],
         [None] * 9,
         ["TEMP", "Temperature test — one question of yours, asked three times at temperature 0 and three at the highest "
                  "the panel allows. Write the six numbers in 'answer_from_tool', separated by ' / '. Did the number move?",
          None, None, None, None, None, None, None]]
sheet(wb, "4 Test",
      ["id", "question", "expected_answer", "where_found", "answer_from_tool", "citation_correct",
       "retrieval_hit", "generation_hit", "note (retrieval / generation / memory, and why)"],
      rows, [7, 50, 22, 26, 30, 12, 12, 12, 36],
      note="Same method as Homework 4. Write each question and find its answer by hand in the Q3 document "
           "BEFORE asking the notebook. An expected answer copied from the tool makes the row worthless.")

# ---- 5 Cost -----------------------------------------------------------------
sheet(wb, "5 Cost",
      ["Item", "Value", "How I got it (or: estimate)"],
      [["Input price per 1M tokens ($)", ref("'1 Model'!B7"), "from tab 1"],
       ["Output price per 1M tokens ($)", ref("'1 Model'!B8"), "from tab 1"],
       ["Tokens sent with one question (AI Studio counter, documents attached)", None, None],
       ["Tokens in one answer", None, None],
       ["Cost of one question ($)", '=IF(OR(B4="",B5=""),"",(B4*B2+B5*B3)/1000000)', "formula"],
       ["Questions behind one analyst note", None, "estimate"],
       ["Cost of one note ($)", '=IF(OR(B6="",B7=""),"",B6*B7)', "formula"],
       ["Analyst's cost per hour ($)", None, "estimate — say from where"],
       ["Hours the notebook saves on one note", None, "estimate"],
       ["Value of one note ($)", '=IF(OR(B9="",B10=""),"",B9*B10)', "formula"],
       ["Notes per month for one desk", None, "estimate"],
       ["Price I would ask a desk per month ($)", None, None],
       ["What is left per month after the tokens ($)", '=IF(OR(B8="",B12="",B13=""),"",B13-B8*B12)', "formula"]],
      [58, 18, 40])

# ---- 6 Score ----------------------------------------------------------------
rows = []
for i in range(1, 6):
    p = i + 3  # row of P{i} in '3 Preview'
    k = i + 1
    rows.append([f"P{i}", ref(f"'3 Preview'!B{p}"), ref(f"'3 Preview'!C{p}"), ref(f"'3 Preview'!D{p}"), None, None,
                 f'=IF(OR(C{k}="",E{k}=""),"",ABS(C{k}-E{k}))',
                 f'=IF(OR(D{k}="",E{k}=""),"",ABS(D{k}-E{k}))',
                 f'=IF(OR(G{k}="",H{k}=""),"",IF(G{k}<H{k},"me",IF(G{k}>H{k},"naive","tie")))'])
sheet(wb, "6 Score",
      ["#", "Prediction", "My number", "Naive: Q3 = Q2", "Actual Q3 2026", "Where found (Q3 document, slide)",
       "My error", "Naive error", "Closer"],
      rows, [6, 34, 14, 14, 14, 34, 12, 12, 10])

# ---- BTC daily --------------------------------------------------------------
ws = wb.create_sheet("BTC daily")
ws.append(["Date", "close"])
for d, v in close.items():
    ws.append([d, float(v)])
for c in ws["A"][1:]:
    c.number_format = "yyyy-mm-dd"
ws.column_dimensions["A"].width = 12

GREY = PatternFill("solid", fgColor="EEEEEE")
for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.fill = GREY

wb.save(DATA / "midterm_sheet.xlsx")
print(f"{len(close)} BTC closes, {close.index.min()} – {close.index.max()}; wrote midterm_sheet.xlsx")
